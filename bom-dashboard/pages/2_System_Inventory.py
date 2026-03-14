"""System Inventory Editor — HEQA.

View (with color coding), add, delete, and edit quantities in Inventory.xlsx.

Row colors:
  🔴 Red    — qty == 0  (out of stock)
  🟡 Yellow — 0 < qty <= threshold  (low stock)
  🟢 Green  — qty > threshold  (in stock)
"""

import datetime
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import pandas as pd
import streamlit as st

from config import (
    DATA_DIR, LOGO_PATH,
    SS_AUTH, SS_CURRENT_USER, SS_IS_ADMIN, SS_DATA, SS_RESULTS,
    INV_FILENAME, INV_SHEET, INV_KEY_COL, INV_QTY_COL,
    TYPE_FILENAME, TYPE_KEY_COL, TYPE_COL,
)
from utils.type_parser import load_types

# ── Page config ────────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="System Inventory — HEQA",
    page_icon="📦",
    layout="wide",
)

st.markdown(
    "<style>[data-testid=\"stSidebarNav\"] { display: none; }</style>",
    unsafe_allow_html=True,
)

with st.sidebar:
    st.markdown("### 🗂️ Navigation")
    if st.button("🏠 Main Dashboard", use_container_width=True, key="nav_main"):
        st.switch_page("app.py")
    if st.button("📤 Upload Files", use_container_width=True, key="nav_site"):
        st.switch_page("pages/1_Site_Inventory.py")
    st.button("📦 System Inventory", use_container_width=True, disabled=True)
    if st.session_state.get(SS_IS_ADMIN):
        if st.button("👥 Users", use_container_width=True, key="nav_users"):
            st.switch_page("pages/3_User_Management.py")
    st.markdown("---")

# ── Auth guard ────────────────────────────────────────────────────────────────
if not st.session_state.get(SS_AUTH):
    st.warning("🔒 Please log in from the main page first.")
    st.stop()

# ── Header ────────────────────────────────────────────────────────────────────
if LOGO_PATH.exists():
    logo_col, title_col = st.columns([1, 3])
    with logo_col:
        st.image(LOGO_PATH.read_bytes(), width=180)
    with title_col:
        st.markdown("# 📦 System Inventory")
        st.caption(
            f"View and edit **{INV_FILENAME}** directly. "
            "Changes affect the procurement calculation immediately."
        )
else:
    st.title("📦 System Inventory")
    st.caption(f"Edit {INV_FILENAME} directly.")

# ── Constants ─────────────────────────────────────────────────────────────────
INV_PATH = DATA_DIR / INV_FILENAME
SS_SYS_INV = "sys_inv_df"
LOW_STOCK_THRESHOLD = 10  # qty <= this → yellow
_BULK_PREFIXES = ("SCR", "SPC", "NUT", "WAS", "WIR")


# ── Helpers ───────────────────────────────────────────────────────────────────
def _load_from_disk() -> pd.DataFrame:
    if not INV_PATH.exists():
        return pd.DataFrame(columns=[INV_KEY_COL, INV_QTY_COL])
    raw = INV_PATH.read_bytes()
    df = pd.read_excel(io.BytesIO(raw), sheet_name=INV_SHEET)
    df.columns = df.columns.str.strip()
    # Normalise key column (header may contain literal \n)
    key_col = INV_KEY_COL
    for col in df.columns:
        if 'מק"ט' in col and 'מדבר' in col:
            key_col = col
            break
    if key_col != INV_KEY_COL:
        df = df.rename(columns={key_col: INV_KEY_COL})
    df[INV_KEY_COL] = df[INV_KEY_COL].astype(str).str.strip()
    df[INV_QTY_COL] = pd.to_numeric(df.get(INV_QTY_COL, 0), errors="coerce").fillna(0).astype(int)
    return df.reset_index(drop=True)


def _save_to_disk(df: pd.DataFrame) -> str | None:
    """Write DataFrame back to Inventory.xlsx (Sheet1). Returns error string or None."""
    try:
        if INV_PATH.exists():
            try:
                import openpyxl
                wb = openpyxl.load_workbook(INV_PATH)
                sheet_names = wb.sheetnames
            except Exception:
                sheet_names = [INV_SHEET]
        else:
            sheet_names = [INV_SHEET]

        if len(sheet_names) <= 1:
            with pd.ExcelWriter(INV_PATH, engine="openpyxl") as writer:
                df.to_excel(writer, sheet_name=INV_SHEET, index=False)
        else:
            from openpyxl import load_workbook
            wb = load_workbook(INV_PATH)
            if INV_SHEET in wb.sheetnames:
                del wb[INV_SHEET]
            ws = wb.create_sheet(INV_SHEET, 0)
            for c_idx, col_name in enumerate(df.columns, 1):
                ws.cell(row=1, column=c_idx, value=col_name)
            for r_idx, row in enumerate(df.itertuples(index=False), 2):
                for c_idx, val in enumerate(row, 1):
                    ws.cell(row=r_idx, column=c_idx, value=val)
            wb.save(INV_PATH)
        return None
    except Exception as exc:
        return str(exc)


def _row_color(qty: int) -> str:
    if qty == 0:
        return "background-color: #ffd6d6"   # red
    if qty <= LOW_STOCK_THRESHOLD:
        return "background-color: #fff3cd"   # yellow
    return "background-color: #d4edda"       # green


def _style_df(df: pd.DataFrame) -> pd.io.formats.style.Styler:
    """Apply per-row background based on Qty."""
    def row_style(row):
        qty = int(row[INV_QTY_COL]) if INV_QTY_COL in row else 0
        color = _row_color(qty)
        return [color] * len(row)

    return df.style.apply(row_style, axis=1)


def _load_types_map() -> dict:
    """Return {pn: type} from Buy & Make.xlsx, or empty dict if not available."""
    type_path = DATA_DIR / TYPE_FILENAME
    if not type_path.exists():
        return {}
    df, err = load_types(type_path.read_bytes())
    if err or df.empty:
        return {}
    return df.set_index(TYPE_KEY_COL)[TYPE_COL].to_dict()


def _invalidate_cache():
    st.session_state.pop(SS_DATA, None)
    st.session_state.pop(SS_RESULTS, None)
    st.cache_data.clear()


# ── Load ──────────────────────────────────────────────────────────────────────
col_reload, _ = st.columns([1, 6])
with col_reload:
    if st.button("🔄 Reload from file", use_container_width=True):
        st.session_state.pop(SS_SYS_INV, None)

if SS_SYS_INV not in st.session_state:
    st.session_state[SS_SYS_INV] = _load_from_disk()

inv_df: pd.DataFrame = st.session_state[SS_SYS_INV]

# ── Merge Type column from Buy & Make.xlsx ────────────────────────────────────
_type_map = _load_types_map()
if _type_map:
    inv_df = inv_df.copy()
    inv_df["Type"] = inv_df[INV_KEY_COL].map(_type_map).fillna("—")
    # Override BULK prefixes
    _bulk_mask = inv_df[INV_KEY_COL].str.upper().str.startswith(_BULK_PREFIXES)
    inv_df.loc[_bulk_mask, "Type"] = "BULK"
else:
    inv_df = inv_df.copy()
    inv_df["Type"] = "—"

if not INV_PATH.exists():
    st.warning(
        f"⚠️ `{INV_FILENAME}` not found in the data folder. "
        "Upload it first via the **Upload Files** page."
    )

# ── KPIs ──────────────────────────────────────────────────────────────────────
total_rows = len(inv_df)
has_qty = INV_QTY_COL in inv_df.columns
out_of_stock = int((inv_df[INV_QTY_COL] == 0).sum()) if has_qty else 0
low_stock = int(((inv_df[INV_QTY_COL] > 0) & (inv_df[INV_QTY_COL] <= LOW_STOCK_THRESHOLD)).sum()) if has_qty else 0
in_stock = int((inv_df[INV_QTY_COL] > LOW_STOCK_THRESHOLD).sum()) if has_qty else 0
total_units = int(inv_df[INV_QTY_COL].sum()) if has_qty else 0

k1, k2, k3, k4, k5 = st.columns(5)
k1.metric("Total Parts", f"{total_rows:,}")
k2.metric("🟢 In Stock", f"{in_stock:,}")
k3.metric("🟡 Low Stock", f"{low_stock:,}")
k4.metric("🔴 Out of Stock", f"{out_of_stock:,}")
k5.metric("Total Units", f"{total_units:,}")

st.caption(
    f"Color thresholds — 🔴 qty = 0 &nbsp;|&nbsp; "
    f"🟡 qty 1–{LOW_STOCK_THRESHOLD} &nbsp;|&nbsp; "
    f"🟢 qty > {LOW_STOCK_THRESHOLD}"
)

st.markdown("---")

# ── Type filter ───────────────────────────────────────────────────────────────
_type_options = ["All"] + sorted([t for t in inv_df["Type"].unique() if t != "—"]) + (["—"] if "—" in inv_df["Type"].values else [])
_sel_type = st.radio("Filter by Type", _type_options, horizontal=True, key="type_filter")

if _sel_type != "All":
    inv_df = inv_df[inv_df["Type"] == _sel_type].reset_index(drop=True)

# ── Tabs: View (colored) | Edit | Add ─────────────────────────────────────────
tab_view, tab_edit, tab_add = st.tabs(["📊 View (colored)", "✏️ Edit / Delete", "➕ Add Part"])

# ──────────────────────────────────────────────────────────────────────────────
# TAB 1 — Colored read-only view
# ──────────────────────────────────────────────────────────────────────────────
with tab_view:
    search_v = st.text_input("🔍 Filter", placeholder="search by P.N or any field…", key="search_view")
    view_df = inv_df.copy()
    if search_v.strip():
        mask = view_df.apply(
            lambda col: col.astype(str).str.contains(search_v.strip(), case=False, na=False)
        ).any(axis=1)
        view_df = view_df[mask].reset_index(drop=True)

    st.caption(f"Showing {len(view_df):,} of {total_rows:,} rows")
    st.dataframe(
        _style_df(view_df),
        use_container_width=True,
        hide_index=True,
    )

    st.markdown("---")
    ex_col, _ = st.columns([1, 4])
    with ex_col:
        st.download_button(
            "⬇️ Export to CSV",
            data=inv_df.to_csv(index=False).encode("utf-8-sig"),
            file_name=f"system_inventory_{datetime.date.today()}.csv",
            mime="text/csv",
            use_container_width=True,
        )

# ──────────────────────────────────────────────────────────────────────────────
# TAB 2 — Edit / Delete
# ──────────────────────────────────────────────────────────────────────────────
with tab_edit:
    search_e = st.text_input("🔍 Filter", placeholder="search by P.N or any field…", key="search_edit")
    edit_source = inv_df.copy()
    if search_e.strip():
        mask = edit_source.apply(
            lambda col: col.astype(str).str.contains(search_e.strip(), case=False, na=False)
        ).any(axis=1)
        edit_source = edit_source[mask].reset_index(drop=True)

    st.caption(f"Showing {len(edit_source):,} of {total_rows:,} rows — edit Qty or any field, then click **Save Changes**.")

    edit_df = edit_source.copy()
    edit_df.insert(0, "Delete", False)

    col_cfg: dict = {
        "Delete":    st.column_config.CheckboxColumn("🗑️", help="Check row(s) to delete"),
        INV_KEY_COL: st.column_config.TextColumn("Heqa P.N", disabled=True),
        INV_QTY_COL: st.column_config.NumberColumn("Qty", min_value=0, step=1),
        "Type":      st.column_config.TextColumn("Type", disabled=True, width="small"),
    }
    for col in edit_df.columns:
        if col not in col_cfg:
            col_cfg[col] = st.column_config.TextColumn(col)

    edited = st.data_editor(
        edit_df,
        column_config=col_cfg,
        disabled=["Heqa P.N", "Type"],
        hide_index=True,
        use_container_width=True,
        num_rows="fixed",
        key="sys_inv_editor",
    )

    btn1, btn2, _ = st.columns([1, 1, 4])

    with btn1:
        if st.button("💾 Save Changes", use_container_width=True, type="primary", key="save_btn"):
            to_delete = set(edited.loc[edited["Delete"] == True, INV_KEY_COL].tolist())

            if search_e.strip():
                # Partial view: merge edits back into full df
                pns_in_view = set(edit_source[INV_KEY_COL].tolist())
                updated_lookup = {
                    str(r[INV_KEY_COL]): r.drop(labels=["Delete"]).to_dict()
                    for _, r in edited.iterrows()
                    if str(r[INV_KEY_COL]) not in to_delete
                }
                new_rows = []
                for _, orig_row in inv_df.iterrows():
                    pn = str(orig_row[INV_KEY_COL])
                    if pn in pns_in_view:
                        if pn in to_delete:
                            continue
                        new_rows.append(updated_lookup.get(pn, orig_row.to_dict()))
                    else:
                        new_rows.append(orig_row.to_dict())
                new_df = pd.DataFrame(new_rows, columns=inv_df.columns)
            else:
                rows = [
                    r.drop(labels=["Delete"]).to_dict()
                    for _, r in edited.iterrows()
                    if not r.get("Delete")
                ]
                new_df = pd.DataFrame(rows, columns=inv_df.columns)

            new_df[INV_QTY_COL] = pd.to_numeric(new_df[INV_QTY_COL], errors="coerce").fillna(0).astype(int)
            new_df = new_df.reset_index(drop=True)
            # Drop derived Type column before saving back to Excel
            save_df = new_df.drop(columns=["Type"], errors="ignore")
            err = _save_to_disk(save_df)
            if err:
                st.error(f"❌ Save failed: {err}")
            else:
                st.session_state[SS_SYS_INV] = save_df
                _invalidate_cache()
                st.success("✅ Saved!")
                st.rerun()

    with btn2:
        to_del_list = edited.loc[edited["Delete"] == True, INV_KEY_COL].tolist()
        lbl = f"🗑️ Delete ({len(to_del_list)})" if to_del_list else "🗑️ Delete Selected"
        if st.button(lbl, use_container_width=True, disabled=len(to_del_list) == 0, key="del_btn"):
            save_df = inv_df.drop(columns=["Type"], errors="ignore")
            save_df = save_df[~save_df[INV_KEY_COL].isin(to_del_list)].reset_index(drop=True)
            err = _save_to_disk(save_df)
            if err:
                st.error(f"❌ Save failed: {err}")
            else:
                st.session_state[SS_SYS_INV] = save_df
                _invalidate_cache()
                st.success(f"✅ Deleted {len(to_del_list)} row(s).")
                st.rerun()

# ──────────────────────────────────────────────────────────────────────────────
# TAB 3 — Add new part
# ──────────────────────────────────────────────────────────────────────────────
with tab_add:
    st.markdown("Fill in the details for a new part and click **Add Part**.")

    _INV_DESC_COL = "תיאור"
    _INV_MFR_COL  = "MFR"
    _INV_MPN_COL  = "MFR P/N for reference"

    with st.form("add_sys_part_form", clear_on_submit=True):
        fa1, fa2, fa3 = st.columns([2, 4, 2])
        with fa1:
            new_pn = st.text_input("Heqa P.N *", placeholder="e.g. 100-0042")
        with fa2:
            new_desc = st.text_input("Description", placeholder="optional")
        with fa3:
            new_qty = st.number_input("Qty", min_value=0, step=1, value=0)

        fb1, fb2 = st.columns(2)
        with fb1:
            new_mfr = st.text_input("MFR Name", placeholder="e.g. Intel")
        with fb2:
            new_mpn = st.text_input("MFR P/N", placeholder="e.g. 10M50DAF484I7G")

        add_btn = st.form_submit_button("➕ Add Part", use_container_width=True)

    if add_btn:
        pn = new_pn.strip()
        if not pn:
            st.error("Heqa P.N is required.")
        elif pn in inv_df[INV_KEY_COL].values:
            st.error(f"❌ `{pn}` already exists — edit its quantity in the **Edit / Delete** tab.")
        else:
            new_row: dict = {col: "" for col in inv_df.columns}
            new_row[INV_KEY_COL] = pn
            new_row[INV_QTY_COL] = int(new_qty)
            if _INV_DESC_COL in inv_df.columns:
                new_row[_INV_DESC_COL] = new_desc.strip()
            if _INV_MFR_COL in inv_df.columns:
                new_row[_INV_MFR_COL] = new_mfr.strip()
            if _INV_MPN_COL in inv_df.columns:
                new_row[_INV_MPN_COL] = new_mpn.strip()
            new_inv = pd.concat([inv_df, pd.DataFrame([new_row])], ignore_index=True)
            err = _save_to_disk(new_inv.drop(columns=["Type"], errors="ignore"))
            if err:
                st.error(f"❌ Could not save: {err}")
            else:
                st.session_state[SS_SYS_INV] = new_inv
                _invalidate_cache()
                qty_color = "🔴" if int(new_qty) == 0 else ("🟡" if int(new_qty) <= LOW_STOCK_THRESHOLD else "🟢")
                st.success(f"✅ Added `{pn}` with qty {int(new_qty)} {qty_color}")
                st.rerun()
